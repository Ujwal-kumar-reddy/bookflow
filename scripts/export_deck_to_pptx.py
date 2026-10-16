import os
import time
from playwright.sync_api import sync_playwright
import pptx
from pptx.util import Inches

def export_slides():
    output_dir = "presentation/rendered_slides"
    os.makedirs(output_dir, exist_ok=True)
    pptx_path = "BookFlow_DevOps_Capstone_Presentation.pptx"

    if os.path.exists(pptx_path):
        os.remove(pptx_path)
        print(f"Removed previous {pptx_path}")

    total_slides = 11
    print(f"Capturing {total_slides} slides from localhost presentation...")

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 1600, 'height': 1000}, device_scale_factor=2)
        
        for slide_num in range(1, total_slides + 1):
            url = f"http://localhost:5050/presentation/index.html?slide={slide_num}&clean=true"
            page.goto(url)
            page.wait_for_selector("#deck")
            # Brief pause to ensure all animations and DOM state settle
            time.sleep(0.3)
            img_path = os.path.join(output_dir, f"slide_{slide_num}.png")
            page.locator("#deck").screenshot(path=img_path)
            print(f"  [OK] Captured Slide {slide_num} / {total_slides} -> {img_path}")

        browser.close()

    print("Building fresh 16:9 PowerPoint presentation from rendered slides...")
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    for slide_num in range(1, total_slides + 1):
        img_path = os.path.join(output_dir, f"slide_{slide_num}.png")
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_path, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))
        print(f"  [OK] Added Slide {slide_num} / {total_slides} to PPTX")

    prs.save(pptx_path)
    print(f"\nSuccessfully generated {pptx_path} with all {total_slides} slides matching localhost presentation exactly!")

if __name__ == "__main__":
    export_slides()
