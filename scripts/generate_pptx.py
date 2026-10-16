import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette: Modern Enterprise White/Indigo
    BG_COLOR = RGBColor(248, 250, 252)       # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)        # #FFFFFF
    CARD_BORDER = RGBColor(226, 232, 240)    # #E2E8F0
    TEXT_MAIN = RGBColor(15, 23, 42)         # #0F172A
    TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B
    PRIMARY_INDIGO = RGBColor(79, 70, 229)   # #4F46E5
    EMERALD = RGBColor(16, 185, 129)         # #10B981
    AMBER = RGBColor(245, 158, 11)           # #F59E0B

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, tag, title, slide_num):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.3))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = PRIMARY_INDIGO

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(10), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_MAIN

        # Slide Number Pill
        num_box = slide.shapes.add_textbox(Inches(11.8), Inches(0.4), Inches(0.8), Inches(0.4))
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = f"{slide_num}/10"
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.font.size = Pt(12)
        p_num.font.bold = True
        p_num.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title="", border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1)

        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), Inches(0.4))
            tf = tb.text_frame
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = PRIMARY_INDIGO
        return card

    # ==========================================================
    # SLIDE 1: Title & Presentation Cover
    # ==========================================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    add_background(s1)

    # Main Center Card
    add_card(s1, Inches(1.5), Inches(1.2), Inches(10.33), Inches(5.1), border_color=PRIMARY_INDIGO)

    # Badge Pill
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(1.8), Inches(3.2), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(238, 242, 255)
    badge.line.fill.background()
    p_b = badge.text_frame.paragraphs[0]
    p_b.text = "DEVOPS CAPSTONE DEFENSE"
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = PRIMARY_INDIGO
    p_b.alignment = PP_ALIGN.CENTER

    # Project Title
    t_box = s1.shapes.add_textbox(Inches(2.0), Inches(2.4), Inches(9.3), Inches(1.2))
    tf = t_box.text_frame
    p = tf.paragraphs[0]
    p.text = "BookFlow: Student Library Lending Tracker"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN

    p2 = tf.add_paragraph()
    p2.text = "Cloud-Native Microservices, CI/CD, Containerization, Kubernetes & Observability"
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED

    # Highlights Row (3 mini cards)
    h_data = [
        ("Production Grade", "FastAPI + React 18 + PostgreSQL"),
        ("Immutable Pipelines", "GitHub Actions + Trivy + GHCR"),
        ("Kubernetes Native", "Helm Charts + Prometheus + Grafana")
    ]
    for i, (hd, sub) in enumerate(h_data):
        x = Inches(2.0 + i * 3.15)
        c = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(4.3), Inches(3.0), Inches(1.4))
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(248, 250, 252)
        c.line.color.rgb = CARD_BORDER
        tf_c = c.text_frame
        p_h = tf_c.paragraphs[0]
        p_h.text = hd
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = PRIMARY_INDIGO
        p_s = tf_c.add_paragraph()
        p_s.text = sub
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = TEXT_MUTED

    # ==========================================================
    # SLIDE 2: Problem Statement & Solution Architecture
    # ==========================================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_background(s2)
    add_header(s2, "Context & Objectives", "Problem Statement & Solution Architecture", 2)

    # Left Card: The Challenge
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), title="The Problem Space")
    tb2_l = s2.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(5.2), Inches(4.3))
    tf2_l = tb2_l.text_frame
    tf2_l.word_wrap = True
    points_l = [
        ("Campus Library Checkout Inefficiency", "Manual book tracking causes lending delays, record mismatches, and untracked overdue loans."),
        ("Fragile Legacy Deployments", "Monolithic setups suffer from host port conflicts, environment drift, and single-point-of-failure downtimes."),
        ("Lack of Operational Visibility", "Zero real-time telemetry into request latency, database query bottlenecks, or service resource saturation.")
    ]
    for idx, (title, desc) in enumerate(points_l):
        p_t = tf2_l.paragraphs[0] if idx == 0 else tf2_l.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = TEXT_MAIN
        p_d = tf2_l.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    # Right Card: The Solution
    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), title="The Cloud-Native Solution")
    tb2_r = s2.shapes.add_textbox(Inches(7.0), Inches(2.3), Inches(5.3), Inches(4.3))
    tf2_r = tb2_r.text_frame
    tf2_r.word_wrap = True
    points_r = [
        ("Containerized Microservices", "FastAPI backend, React 18 frontend, and PostgreSQL containerized with non-root security."),
        ("Strict Automated Quality Gates", "Automated Pytest testing (16/16) and Trivy vulnerability gate (0 CVEs) in GitHub Actions."),
        ("Declarative Kubernetes Orchestration", "Helm package management, 2-replica High Availability, NGINX Ingress, and persistent storage."),
        ("Continuous Observability", "Prometheus metrics harvesting and Grafana dashboard for live request rates and memory profiling.")
    ]
    for idx, (title, desc) in enumerate(points_r):
        p_t = tf2_r.paragraphs[0] if idx == 0 else tf2_r.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = EMERALD
        p_d = tf2_r.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    # ==========================================================
    # SLIDE 3: Technology Stack & Toolchain
    # ==========================================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_background(s3)
    add_header(s3, "Architecture Stack", "Technology Stack & Toolchain", 3)

    tech_pillars = [
        ("Application Core", ["FastAPI (Python 3.11/3.13)", "React 18 & Vite 5", "PostgreSQL 16 Engine", "Alembic Migrations", "SQLAlchemy ORM"]),
        ("Testing & Quality", ["Pytest 8.3 Suite (16 Tests)", "SQLite In-Memory Mock", "Pre-Push Secret Scanning", "Vite Production Bundler", "Docker Compose Local Stack"]),
        ("CI/CD & Security", ["GitHub Actions Automation", "Trivy CVE Container Scanner", "GitHub Container Registry", "Immutable Git SHA Tagging", "Non-Root UID Hardening"]),
        ("K8s & Observability", ["Kubernetes v1.31 Runtime", "Helm 3 Declarative Charts", "NGINX Ingress Controller", "Prometheus & ServiceMonitor", "Grafana Live Dashboards & HPA"])
    ]
    for idx, (pillar, items) in enumerate(tech_pillars):
        x = Inches(0.8 + idx * 2.95)
        add_card(s3, x, Inches(1.6), Inches(2.8), Inches(5.2), title=pillar)
        tb = s3.shapes.add_textbox(x + Inches(0.15), Inches(2.3), Inches(2.5), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        for i_idx, itm in enumerate(items):
            p = tf.paragraphs[0] if i_idx == 0 else tf.add_paragraph()
            p.text = f"▸ {itm}"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_MAIN

    # ==========================================================
    # SLIDE 4: Core Application & Enterprise UI
    # ==========================================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_background(s4)
    add_header(s4, "Frontend & API", "Enterprise Application & Circulation Desk UI", 4)

    # Left: Explanation
    add_card(s4, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2), title="Circulation Desk Features")
    tb4 = s4.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(4.1), Inches(4.3))
    tf4 = tb4.text_frame
    tf4.word_wrap = True
    ui_points = [
        ("White Enterprise Palette", "High-contrast clean layout inspired by modern SaaS standards (#F8FAFC canvas, crisp #FFFFFF cards, indigo accents)."),
        ("Live Library KPI Deck", "Instant count of Total Books (7), Available on Shelves (4), Active Loans (3), and Overdue Books (0)."),
        ("Search & Dynamic Category Filtering", "Fast client-side search across title, author, and borrower with multi-category chips."),
        ("Interactive Loan & Return Workflows", "Modals for lending books to students (with roll ID & loan period) and catalogue addition.")
    ]
    for idx, (title, desc) in enumerate(ui_points):
        p_t = tf4.paragraphs[0] if idx == 0 else tf4.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = TEXT_MAIN
        p_d = tf4.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_MUTED

    # Right: Embedded Screenshot
    img_card = add_card(s4, Inches(5.5), Inches(1.6), Inches(7.0), Inches(5.2))
    img_path = r"evidence\m1-app-browser.png"
    if os.path.exists(img_path):
        s4.shapes.add_picture(img_path, Inches(5.6), Inches(1.8), Inches(6.8), Inches(4.8))

    # ==========================================================
    # SLIDE 5: Quality Assurance & Security Hardening
    # ==========================================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_background(s5)
    add_header(s5, "Security & Verification", "Quality Assurance & Student Loan Workflow", 5)

    add_card(s5, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2), title="Automated Verification & Controls")
    tb5 = s5.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(4.1), Inches(4.3))
    tf5 = tb5.text_frame
    tf5.word_wrap = True
    qa_points = [
        ("16/16 Pytest Test Suite", "Full coverage for CRUD, health, ready, and metrics routes in 1.32s."),
        ("Zero Database Pollution", "In-memory SQLite with StaticPool ensures 0 leftover test records."),
        ("Non-Root Security", "Backend runs as UID 10001 (appuser); frontend runs as UID 101 (nginx)."),
        ("Interactive Checkout Modal", "Validates student roll IDs and dynamic loan period calculations.")
    ]
    for idx, (title, desc) in enumerate(qa_points):
        p_t = tf5.paragraphs[0] if idx == 0 else tf5.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = EMERALD
        p_d = tf5.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    add_card(s5, Inches(5.5), Inches(1.6), Inches(7.0), Inches(5.2))
    img5_path = r"evidence\m1-loan-modal.png"
    if os.path.exists(img5_path):
        s5.shapes.add_picture(img5_path, Inches(5.7), Inches(1.85), Inches(6.6), Inches(4.7))

    # ==========================================================
    # SLIDE 6: CI/CD Pipeline with GitHub Actions
    # ==========================================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    add_background(s6)
    add_header(s6, "Automated Pipeline", "Continuous Integration & Delivery Pipeline", 6)

    pipeline_stages = [
        ("Stage 1: Test Gate", "Trigger: Push to 'main'", [
            "Checkout repository code",
            "Set up Python 3.11 environment",
            "Execute 16/16 Pytest tests (blocking)",
            "Set up Node.js 20 environment",
            "Run npm ci && npm run build (passed)"
        ]),
        ("Stage 2: Vulnerability Gate", "Security First", [
            "Build Docker images for backend & frontend",
            "Run Aquasecurity Trivy CVE scanner",
            "Strict gate: exit-code 1 on HIGH or CRITICAL",
            "Scan Result: 0 High, 0 Critical vulnerabilities",
            "Guarantees clean container posture"
        ]),
        ("Stage 3: GHCR Publishing", "Immutable Artifacts", [
            "Authenticate with GitHub Container Registry",
            "Tag images with exact Git commit SHA",
            "Push backend: <sha> to GHCR",
            "Push frontend: <sha> to GHCR",
            "Deterministic image deployment traceability"
        ])
    ]
    for idx, (stg, sub, items) in enumerate(pipeline_stages):
        x = Inches(0.8 + idx * 3.95)
        add_card(s6, x, Inches(1.6), Inches(3.8), Inches(5.2), title=stg)
        tb = s6.shapes.add_textbox(x + Inches(0.2), Inches(2.3), Inches(3.4), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p_sub = tf.paragraphs[0]
        p_sub.text = sub.upper()
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = TEXT_MUTED

        for itm in items:
            p = tf.add_paragraph()
            p.text = f"▸ {itm}"
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_MAIN

    # ==========================================================
    # SLIDE 7: Kubernetes & Declarative Helm Deployment
    # ==========================================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    add_background(s7)
    add_header(s7, "Orchestration & Helm", "Kubernetes & Declarative Helm Architecture", 7)

    # Left: Explanation
    add_card(s7, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2), title="Production Orchestration")
    tb7 = s7.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(4.1), Inches(4.3))
    tf7 = tb7.text_frame
    tf7.word_wrap = True
    k8s_points = [
        ("Custom Helm Chart", "Modular chart in helm/bookflow with linted values (0 errors)."),
        ("High Availability (5/5)", "2 backend replicas + 2 frontend replicas + 1 PostgreSQL stateful pod."),
        ("Traffic Ingress Controller", "NGINX Ingress routing root (/) to frontend and (/api, /metrics) to backend."),
        ("Probes & Self-Healing", "Liveness and readiness HTTP probes preventing traffic to unready pods."),
        ("Persistent Storage", "Kubernetes PVC mount ensuring PostgreSQL data durability across restarts.")
    ]
    for idx, (title, desc) in enumerate(k8s_points):
        p_t = tf7.paragraphs[0] if idx == 0 else tf7.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = TEXT_MAIN
        p_d = tf7.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    # Right: Embedded Screenshot
    add_card(s7, Inches(5.5), Inches(1.6), Inches(7.0), Inches(5.2))
    img7_path = r"evidence\m8-ingress-browser.png"
    if os.path.exists(img7_path):
        s7.shapes.add_picture(img7_path, Inches(5.7), Inches(1.85), Inches(6.6), Inches(4.7))

    # ==========================================================
    # SLIDE 8: Kubernetes Troubleshooting Labs
    # ==========================================================
    s8 = prs.slides.add_slide(blank_slide_layout)
    add_background(s8)
    add_header(s8, "Reliability Engineering", "Kubernetes Hands-On Troubleshooting Labs", 8)

    # Lab 1
    add_card(s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), title="Lab 1: ImagePullBackOff Diagnosis")
    tb8_1 = s8.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(5.2), Inches(4.3))
    tf8_1 = tb8_1.text_frame
    tf8_1.word_wrap = True
    lab1_pts = [
        ("Fault Injected", "Deployed broken manifest with non-existent tag :v999.999-broken."),
        ("Observed Failure", "Pod entered ErrImagePull and escalated to ImagePullBackOff (0/1 ready)."),
        ("Diagnosis Protocol", "Inspected Kubelet events via kubectl describe pod, confirming 404 tag mismatch."),
        ("Remediation", "Patched deployment with valid image tag (ghcr.io/...:local-verified). Pod recovered to Running."),
        ("Audit Evidence", "Captured full CLI diagnosis log in evidence/m8-troubleshooting-image.txt.")
    ]
    for idx, (title, desc) in enumerate(lab1_pts):
        p_t = tf8_1.paragraphs[0] if idx == 0 else tf8_1.add_paragraph()
        p_t.text = f"▸ {title}: {desc}"
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = TEXT_MAIN

    # Lab 2
    add_card(s8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), title="Lab 2: Missing Endpoints Diagnosis")
    tb8_2 = s8.shapes.add_textbox(Inches(7.0), Inches(2.3), Inches(5.3), Inches(4.3))
    tf8_2 = tb8_2.text_frame
    tf8_2.word_wrap = True
    lab2_pts = [
        ("Fault Injected", "Configured Service with typo in selector: app=bookflow-be-wrong."),
        ("Observed Failure", "kubectl get endpoints returned <none>; curl requests received 502 Bad Gateway."),
        ("Diagnosis Protocol", "Compared pod labels with Service spec.selector to isolate the label discrepancy."),
        ("Remediation", "Aligned Service selector to match pod labels; endpoints immediately populated 2 IPs."),
        ("Audit Evidence", "Captured full CLI diagnosis log in evidence/m8-troubleshooting-service.txt.")
    ]
    for idx, (title, desc) in enumerate(lab2_pts):
        p_t = tf8_2.paragraphs[0] if idx == 0 else tf8_2.add_paragraph()
        p_t.text = f"▸ {title}: {desc}"
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = TEXT_MAIN

    # ==========================================================
    # SLIDE 9: Full-Stack Observability & Autoscaling
    # ==========================================================
    s9 = prs.slides.add_slide(blank_slide_layout)
    add_background(s9)
    add_header(s9, "Telemetry & Metrics", "Full-Stack Observability & Autoscaling (HPA)", 9)

    # Left: Explanation
    add_card(s9, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2), title="Observability Architecture")
    tb9 = s9.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(4.1), Inches(4.3))
    tf9 = tb9.text_frame
    tf9.word_wrap = True
    obs_points = [
        ("Prometheus Instrumentator", "FastAPI middleware exposing http_requests_total, response latency histograms, and process metrics."),
        ("Active Scrape Target", "Prometheus operator scraping the backend service every 15s (Target Health: UP)."),
        ("Live 4-Panel Grafana Dashboard", "Custom dashboard monitoring Request Rate (QPS), Request Latency (p95), 5xx Error Rate, and Process Memory."),
        ("Horizontal Pod Autoscaler (HPA)", "Kubernetes HPA configured with 60% CPU target utilization (scales from 2 to 5 pods under spike load).")
    ]
    for idx, (title, desc) in enumerate(obs_points):
        p_t = tf9.paragraphs[0] if idx == 0 else tf9.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = TEXT_MAIN
        p_d = tf9.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    # Right: Embedded Screenshot (Bounded)
    add_card(s9, Inches(5.5), Inches(1.6), Inches(7.0), Inches(5.2))
    img9_path = r"evidence\m9-grafana-dashboard.png"
    if os.path.exists(img9_path):
        s9.shapes.add_picture(img9_path, Inches(5.7), Inches(1.85), Inches(6.6), Inches(4.7))

    # ==========================================================
    # SLIDE 10: Capstone Summary & Production Takeaways (2x2 Grid)
    # ==========================================================
    s10 = prs.slides.add_slide(blank_slide_layout)
    add_background(s10)
    add_header(s10, "Summary & Impact", "Capstone Deliverables & DevOps Mastery", 10)

    summary_boxes = [
        ("100% Verified Local Stack", "Docker Desktop Kubernetes running 5/5 healthy pods, NGINX Ingress on port 8088, and persistent PostgreSQL storage surviving pod restarts."),
        ("Zero Secret Leaks & $0 Cloud Costs", "Automated pre-push scanner passed with 0 secret patterns; Terraform IaC statically validated without unexpected AWS charges ($0.00 spent)."),
        ("Production CI/CD Automation", "GitHub Actions pipeline with 16/16 Pytest blocking gate, Trivy CVE security checks (0 CVEs), and immutable GHCR image tags."),
        ("Observability & 37 Evidence Artifacts", "Prometheus metrics scraping with target status UP, live 4-panel Grafana dashboard, HPA dynamic scaling, and 37 comprehensive evidence files.")
    ]
    for idx, (head, desc) in enumerate(summary_boxes):
        col = idx % 2
        row = idx // 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.6 + row * 2.7)
        c = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(2.5))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER
        tf = c.text_frame
        p = tf.paragraphs[0]
        p.text = f"✔  {head}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_INDIGO if col == 1 else EMERALD
        p_d = tf.add_paragraph()
        p_d.text = f"     {desc}"
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MUTED

    # Save Presentation
    output_path = "BookFlow_DevOps_Capstone_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_presentation()
